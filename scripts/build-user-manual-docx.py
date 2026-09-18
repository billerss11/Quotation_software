"""Build searchable Word user guides from maintained Markdown and real screenshots.

Uses python-docx from the document authoring environment, not the app's dependencies.
First run export-user-manual-content.mjs; see docs/user-manual-screenshots.md.
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


# Put the picture immediately after the relevant instructions, before the next heading.
FIGURES = {
    '3.1': ['overview'], '4.1': ['settings'], '5.1': ['company'],
    '6.1': ['customers'], '7.1': ['new-menu'], '7.3': ['save'],
    '7.4': ['open-menu'], '7.5': ['import-quotation'], '9.1': ['quote-info'],
    '10.2': ['parties'], '12.1': ['add-item'], '12.4': ['edit-item'],
    '13.1': ['add-child'], '13.2': ['hierarchy'], '17.2': ['pricing'],
    '19.2': ['fx'], '23.1': ['preview-document'], '23.2': ['preview-options'],
    '24.1': ['analysis'], '25.1': ['receipt-open'], '25.3': ['receipt-dialog'],
    '25.4': ['receipt-selection'], '25.5': ['receipt-items'],
    '26.3': ['import-dialog'],
}

# Crop metadata in the DOCX keeps the original PNG intact and focuses on the controls.
# Rectangles are x, y, width, height in source pixels. No screenshot pixels are edited.
CONTEXT_CROPS = {
    'overview': (0, 0, 1440, 960),
    'settings': (85, 70, 1330, 850),
    'new-menu': (960, 0, 480, 380),
    'open-menu': (960, 0, 480, 380),
    'import-quotation': (960, 0, 480, 380),
    'save': (735, 0, 705, 130),
    'receipt-open': (965, 0, 475, 135),
    'quote-info': (1005, 258, 400, 652),
    'parties': (1005, 256, 400, 585),
    'pricing': (1005, 265, 400, 670),
    'fx': (1005, 160, 400, 400),
    'preview-options': (1005, 430, 430, 485),
    'preview-document': (285, 16, 875, 917),
    'analysis': (80, 88, 1340, 765),
    'receipt-dialog': (20, 170, 550, 515),
    'receipt-selection': (25, 452, 535, 405),
    'receipt-items': (23, 490, 542, 370),
    'import-dialog': (262, 225, 916, 525),
}


def element(tag, **attrs):
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(qn(key), str(value))
    return node


def native_link(paragraph, text, target, internal=False):
    link = OxmlElement('w:hyperlink')
    if internal:
        link.set(qn('w:anchor'), target)
    else:
        relation = paragraph.part.relate_to(
            target, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
            is_external=True)
        link.set(qn('r:id'), relation)
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    props.append(element('w:color', **{'w:val': '175D69'}))
    props.append(element('w:u', **{'w:val': 'single'}))
    run.append(props)
    text_node = OxmlElement('w:t')
    text_node.text = text
    run.append(text_node)
    link.append(run)
    paragraph._p.append(link)


def inline(paragraph, tokens, bold=False, italic=False):
    for token in tokens or []:
        kind = token['type']
        if kind in ('strong', 'em'):
            inline(paragraph, token.get('tokens'), bold or kind == 'strong', italic or kind == 'em')
        elif kind == 'link':
            label = token.get('text', token.get('href', ''))
            href = token.get('href', '')
            if href.startswith(('https://', 'http://', 'mailto:')):
                native_link(paragraph, label, href)
            else:
                paragraph.add_run(label).bold = bold
        elif kind == 'br':
            paragraph.add_run().add_break()
        elif token.get('tokens'):
            inline(paragraph, token['tokens'], bold, italic)
        else:
            run = paragraph.add_run(token.get('text', token.get('raw', '')))
            run.bold = bold
            run.italic = italic
            if kind == 'codespan':
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)


def bookmark(paragraph, name, index):
    paragraph._p.insert(0, element('w:bookmarkStart', **{'w:id': index, 'w:name': name}))
    paragraph._p.append(element('w:bookmarkEnd', **{'w:id': index}))


class Guide:
    def __init__(self, data, locale):
        self.data, self.locale = data, locale
        self.zh = locale == 'zh-CN'
        self.content = data['locales'][locale]
        self.root = Path(data['root'])
        self.doc = Document()
        self.figure_count = 0
        self.bookmark_count = 0
        self.table_count = 0
        self.section_count = 0
        self.steps = {step['id']: step for topic in data['sections'] for step in topic['steps']}
        self.setup_styles()

    def setup_styles(self):
        section = self.doc.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin, section.bottom_margin = Inches(.65), Inches(.65)
        section.left_margin, section.right_margin = Inches(.7), Inches(.7)
        section.footer_distance = Inches(.28)
        for grid in list(section._sectPr.findall(qn('w:docGrid'))):
            section._sectPr.remove(grid)
        styles = self.doc.styles
        for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3',
                     'Caption', 'List Bullet', 'List Bullet 2', 'List Number', 'List Number 2', 'Footer'):
            style = styles[name]
            style.font.name = 'Arial'
            style.font.color.rgb = RGBColor(0, 0, 0)
            paragraph_props = style.element.get_or_add_pPr()
            for border in list(paragraph_props.findall(qn('w:pBdr'))):
                paragraph_props.remove(border)
            paragraph_props.append(element('w:snapToGrid', **{'w:val': 'false'}))
            style.paragraph_format.page_break_before = False
            rpr = style.element.get_or_add_rPr()
            rfonts = rpr.find(qn('w:rFonts'))
            if rfonts is None:
                rfonts = OxmlElement('w:rFonts')
                rpr.append(rfonts)
            for attribute, value in [('ascii', 'Arial'), ('hAnsi', 'Arial'), ('eastAsia', 'Microsoft YaHei')]:
                rfonts.set(qn('w:' + attribute), value)
            for attribute in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
                rfonts.attrib.pop(qn('w:' + attribute), None)
            for c in rpr.findall(qn('w:color')):
                c.attrib.pop(qn('w:themeColor'), None)
            lang = element('w:lang', **{'w:val': self.locale, 'w:eastAsia': 'zh-CN'})
            rpr.append(lang)
        normal = styles['Normal']
        normal.font.size = Pt(11)
        normal.paragraph_format.line_spacing = Pt(16) if self.zh else 1.15
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.widow_control = True
        for name, size, before, after in [('Title', 25, 0, 12), ('Heading 1', 16, 18, 9),
                                           ('Heading 2', 12, 12, 7), ('Heading 3', 11, 10, 5)]:
            style = styles[name]
            style.font.size = Pt(size)
            style.font.bold = name != 'Title'
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.keep_together = True
            style.paragraph_format.line_spacing = 1.05
        styles['Subtitle'].font.size = Pt(13)
        styles['Subtitle'].font.italic = False
        styles['Caption'].font.size = Pt(9)
        styles['Caption'].font.italic = False
        styles['Caption'].font.bold = False
        styles['Caption'].paragraph_format.space_after = Pt(12)
        styles['Footer'].font.size = Pt(9)
        styles['Footer'].paragraph_format.line_spacing = Pt(12)
        styles['Footer'].paragraph_format.space_after = Pt(0)
        for name in ('List Bullet', 'List Number', 'List Bullet 2', 'List Number 2'):
            styles[name].font.size = Pt(11)
            styles[name].paragraph_format.space_after = Pt(4)
        self.doc.core_properties.title = '报价软件用户指南' if self.zh else 'Quotation Software User Guide'
        self.doc.core_properties.subject = '图文操作说明和完整功能参考' if self.zh else 'Illustrated instructions and complete feature reference'
        self.doc.core_properties.author = 'Quotation Software'
        self.doc.core_properties.keywords = 'quotation, user guide, 报价, 用户指南, Excel, tax, goods receipt'
        self.doc.core_properties.language = self.locale
        settings = self.doc.settings.element
        settings.append(element('w:doNotAutoCompressPictures'))
        settings.append(element('w:updateFields', **{'w:val': 'true'}))
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run('第 ' if self.zh else 'Page ').font.size = Pt(9)
        footer._p.append(element('w:fldSimple', **{'w:instr': 'PAGE'}))
        footer.add_run(' 页' if self.zh else '').font.size = Pt(9)

    def front_matter(self):
        self.doc.add_paragraph(self.doc.core_properties.title, 'Title')
        self.doc.add_paragraph('带截图的操作说明与完整功能参考' if self.zh else 'Illustrated instructions and complete feature reference', 'Subtitle')
        self.doc.add_paragraph('2026 年 9 月 18 日  |  简体中文版' if self.zh else '18 September 2026  |  English edition')
        self.doc.add_paragraph(
            '本指南帮助您创建、检查、保存和输出报价，以及管理客户、公司资料和收货单。操作说明、规则、表格和故障排查均为可搜索、可复制的 Word 文本；截图用于说明界面位置。'
            if self.zh else
            'Use this guide to create, check, save, and output quotations, manage reusable records, and prepare goods receipts. Instructions, rules, tables, and troubleshooting are searchable, copyable Word text. Screenshots show where the controls are.')
        self.doc.add_paragraph('如何查找答案' if self.zh else 'Finding an answer', 'Heading 2')
        for sentence in ([
            '按 Ctrl+F 搜索按钮名称、问题或关键词，例如“混合税率”“markup_override”“收货单”。',
            '使用 Word 导航窗格中的“标题”，或点击目录中的章节链接，跳转到对应内容。',
            '查看截图细节时放大 Word 视图。图片保留原始分辨率；部分图片裁去无关区域以便阅读。',
            '桌面版和浏览器版共用编辑器，但文件按钮不同。截图来自浏览器版；第 2 节说明“下载”和“保存”、“打印”和“导出 PDF”的对应关系。',
        ] if self.zh else [
            'Press Ctrl+F to search for a button, problem, or keyword such as mixed tax, markup_override, or goods receipt.',
            'Use Headings in Word’s Navigation Pane, or follow the linked contents, to jump directly to a section.',
            'Zoom the Word view to inspect screenshot details. Original image resolution is retained; some figures crop unrelated areas for readability.',
            'Desktop and browser editions share the editor but use different file buttons. Screenshots show the browser edition; section 2 maps Download to Save and Print to Export PDF.',
        ]):
            self.doc.add_paragraph(sentence, 'List Bullet')
        self.doc.add_paragraph('使用 AI 提问' if self.zh else 'Asking an AI about this guide', 'Heading 2')
        self.doc.add_paragraph(
            '将此 DOCX 上传到支持文档读取的 AI 工具。可提问：“请根据这份用户指南，说明如何完成［任务］。列出具体按钮、桌面版与浏览器版的区别，并引用章节编号。指南没有说明的部分请明确指出。”'
            if self.zh else
            'Upload this DOCX to an AI tool that accepts documents. Try: “Using this user guide, explain how to [task]. Give the exact buttons, distinguish desktop and browser behavior, and cite the section number. Clearly identify anything the guide does not explain.”')
        self.doc.add_paragraph(
            '截图示例为 Q-SLB-INS-20260913-01。示例包含未完整项目，用于说明操作；实际发送前应核对本报价的数量、价格、税率和条款。'
            if self.zh else
            'The screenshots use quotation Q-SLB-INS-20260913-01. It contains incomplete items and illustrates the workflow. Review the quantities, prices, tax rates, and terms of your own quotation before sending it.')
        self.doc.add_page_break()
        self.doc.add_paragraph('目录' if self.zh else 'Contents', 'Title')
        self.doc.add_paragraph('点击章节标题即可跳转；Ctrl+F 可搜索全文。' if self.zh else 'Follow a section link to jump to it, or press Ctrl+F to search the entire guide.')
        heading_index = 0
        for token in self.content['reference']:
            if token['type'] != 'heading':
                continue
            heading_index += 1
            if token['depth'] != 2:
                continue
            toc = self.doc.add_paragraph()
            toc.paragraph_format.space_after = Pt(3)
            toc.paragraph_format.line_spacing = Pt(14)
            toc.paragraph_format.keep_together = True
            toc.paragraph_format.tab_stops.add_tab_stop(Inches(7.1), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            name = f'manual_section_{heading_index}'
            native_link(toc, re.sub(r'[*`]', '', token['text']), name, internal=True)
            toc.add_run('\t')
            toc._p.append(element('w:fldSimple', **{'w:instr': f'PAGEREF {name} \\h'}))

    def table(self, token):
        header = token['header']
        columns = len(header)
        table = self.doc.add_table(rows=1, cols=columns)
        table.autofit = False
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        fractions = {2: [.28, .72], 3: [.24, .38, .38], 4: [.18, .26, .24, .32]}.get(columns, [1/columns]*columns)
        widths = [Inches(7.1 * fraction) for fraction in fractions]
        for index, width in enumerate(widths):
            table.columns[index].width = width
        borders = OxmlElement('w:tblBorders')
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            borders.append(element('w:' + side, **{'w:val': 'single', 'w:sz': 4, 'w:color': 'D9D9D9'}))
        table._tbl.tblPr.append(borders)
        for row_index, cells in enumerate([header, *token['rows']]):
            row = table.rows[0] if row_index == 0 else table.add_row()
            props = row._tr.get_or_add_trPr()
            props.append(element('w:cantSplit'))
            if row_index == 0:
                props.append(element('w:tblHeader'))
            for index, cell_data in enumerate(cells):
                cell = row.cells[index]
                cell.width = widths[index]
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = cell._tc.get_or_add_tcPr()
                tcpr.append(element('w:shd', **{'w:fill': '254E5A' if row_index == 0 else ('F3F6F7' if row_index % 2 == 0 else 'FFFFFF'), 'w:val': 'clear'}))
                margins = OxmlElement('w:tcMar')
                for side in ('top', 'bottom', 'left', 'right'):
                    margins.append(element('w:' + side, **{'w:w': 90, 'w:type': 'dxa'}))
                tcpr.append(margins)
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.05
                if row_index == 0:
                    paragraph.paragraph_format.keep_with_next = True
                inline(paragraph, cell_data.get('tokens'))
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    if row_index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
        self.table_count += 1
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def block(self, token, level=0):
        kind = token['type']
        if kind in ('space', 'hr'):
            return
        if kind == 'table':
            self.table(token)
        elif kind == 'list':
            numbering_id = None
            if token['ordered']:
                numbering = self.doc.part.numbering_part.element
                style_name = 'List Number 2' if level else 'List Number'
                style_id = self.doc.styles[style_name].element.pPr.numPr.numId.val
                base = numbering.num_having_numId(style_id).abstractNumId.val
                number = numbering.add_num(base)
                override = number.add_lvlOverride(ilvl=0)
                override.add_startOverride(token.get('start', 1) or 1)
                numbering_id = number.numId
            for item in token['items']:
                first = True
                for child in item.get('tokens', []):
                    if child['type'] == 'list':
                        self.block(child, level + 1)
                    elif child['type'] in ('text', 'paragraph'):
                        style = ('List Number' if token['ordered'] else 'List Bullet') + (' 2' if level else '')
                        paragraph = self.doc.add_paragraph(style=style if first else 'Normal')
                        if first and numbering_id is not None:
                            number_props = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                            number_props.get_or_add_ilvl().val = 0
                            number_props.get_or_add_numId().val = numbering_id
                        if not first:
                            paragraph.paragraph_format.left_indent = Inches(.3 + .2 * level)
                        inline(paragraph, child.get('tokens') or [{'type': 'text', 'text': child.get('text', '')}])
                        first = False
                    else:
                        self.block(child, level)
        elif kind == 'code':
            for line in token['text'].splitlines():
                paragraph = self.doc.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.line_spacing = 1
                run = paragraph.add_run(line)
                run.font.name = 'Consolas'
                run.font.size = Pt(8.5)
        elif kind == 'blockquote':
            for child in token.get('tokens', []):
                self.block(child, level)
        elif kind in ('paragraph', 'text'):
            paragraph = self.doc.add_paragraph()
            inline(paragraph, token.get('tokens') or [{'type': 'text', 'text': token.get('text', '')}])
        elif kind == 'html':
            text = re.sub('<[^>]+>', '', token.get('text', token.get('raw', ''))).strip()
            if text:
                self.doc.add_paragraph(text)
        else:
            raise ValueError(f'Unsupported Markdown token: {kind}')

    def figure(self, image_id):
        step = self.steps[image_id]
        message = self.content['messages']['steps'][image_id]
        focus = step.get('focusZh', step['focus']) if self.zh else step['focus']
        if image_id in CONTEXT_CROPS:
            x, y, width, height = CONTEXT_CROPS[image_id]
        else:
            x, y = focus['x'] * 14.4 - 24, focus['y'] * 9.6 - 24
            width, height = focus['width'] * 14.4 + 48, focus['height'] * 9.6 + 48
            if image_id == 'add-item':
                x, y, width, height = max(80, x - 160), max(78, y - 65), width + 320, height + 135
            elif image_id == 'add-child':
                x, y, width, height = 100, max(260, y - 270), 895, 345
        x, y = max(0, x), max(0, y)
        width, height = min(width, 1440 - x), min(height, 960 - y)
        max_height = 5.8 if height > width else 4.65
        factor = min(7.1 / width, max_height / height, 1 / 96)
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(9)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.line_spacing = 1
        shape = paragraph.add_run().add_picture(
            str(self.root / 'docs' / 'assets' / 'user-manual' / self.locale / f'{image_id}.png'),
            width=Inches(width * factor), height=Inches(height * factor))
        blip = shape._inline.graphic.graphicData.pic.blipFill
        src = OxmlElement('a:srcRect')
        for key, value in [('l', x / 1440), ('t', y / 960), ('r', (1440 - x - width) / 1440), ('b', (960 - y - height) / 960)]:
            src.set(key, str(round(value * 100000)))
        blip.insert(1, src)
        shape._inline.docPr.set('descr', message['alt'])
        shape._inline.docPr.set('title', message['title'])
        self.figure_count += 1
        caption = self.doc.add_paragraph(style='Caption')
        caption.paragraph_format.keep_together = True
        caption.add_run(f'图 {self.figure_count}  ' if self.zh else f'Figure {self.figure_count}  ').bold = True
        caption.add_run(message['title'])
        if image_id == 'hierarchy':
            caption.add_run('。横向滚动子项表可查看其余价格和税额列。' if self.zh else '. Scroll the child table horizontally to see the remaining price and tax columns.')
        elif image_id == 'receipt-dialog':
            caption.add_run('。图示为上方字段；向下滚动左侧面板可填写其余信息。' if self.zh else '. Upper fields are shown; scroll the left panel for the remaining details.')

    def build(self, output):
        self.front_matter()
        current_section = None
        for token in self.content['reference']:
            if token['type'] != 'heading':
                self.block(token)
                continue
            for image_id in FIGURES.get(current_section, []):
                self.figure(image_id)
            text = token['text']
            match = re.match(r'^(\d+(?:\.\d+)*)', text)
            current_section = match.group(1) if match else None
            depth = min(token['depth'] - 1, 3)
            heading_text = re.sub(r'[*`]', '', text)
            paragraph = self.doc.add_paragraph(heading_text, f'Heading {depth}')
            if depth == 1:
                self.section_count += 1
                if self.section_count == 1:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.page_break_before = True
            self.bookmark_count += 1
            bookmark(paragraph, f'manual_section_{self.bookmark_count}', self.bookmark_count)
        for image_id in FIGURES.get(current_section, []):
            self.figure(image_id)
        if self.figure_count != 24 or self.section_count != 31:
            raise ValueError(f'Coverage mismatch: {self.figure_count} figures, {self.section_count} sections')
        output.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(output)
        return {'locale': self.locale, 'sections': self.section_count, 'headings': self.bookmark_count,
                'figures': self.figure_count, 'tables': self.table_count, 'output': str(output)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('content_json', type=Path)
    parser.add_argument('--output-dir', type=Path, default=Path('docs/user-guides'))
    parser.add_argument('--locale', choices=['en-US', 'zh-CN'])
    args = parser.parse_args()
    data = json.loads(args.content_json.read_text(encoding='utf8'))
    reports = []
    for locale in ([args.locale] if args.locale else data['locales']):
        reports.append(Guide(data, locale).build(args.output_dir / f'Quotation-Software-User-Guide-{locale}.docx'))
    print(json.dumps(reports, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
